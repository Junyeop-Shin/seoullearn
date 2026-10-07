"""평가셋 검사 + 검수용 엑셀 내보내기.

- 근거(evidence)가 실제 문서·청크 페이지를 가리키는지
- 필수 필드, answerable/route 값
- 유형별 분포 출력
- eval/eval_set_review.xlsx 생성 (검수 칸 포함)
"""
import json
import re
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[1]
ANSWERABLE = {"yes", "partial", "no"}
ROUTES = {None, "운영사무국", "운영사무국(즉시)", "이음단", "학습지원센터"}
MARK = re.compile(r"<!-- page: (?P<page>[^|]+?) \| type:")


def chunk_pages():
    pages = {}
    for f in list((ROOT / "data/manual").glob("*.md")) + list((ROOT / "data/notice").glob("*.md")):
        text = f.read_text(encoding="utf-8")
        did = re.search(r"^doc_id: (.+)$", text, re.M).group(1).strip()
        pages[did] = {m["page"].strip() for m in MARK.finditer(text)}
    return pages


def check_side(item, side, data, pages, errors):
    if data.get("answerable") not in ANSWERABLE:
        errors.append(f"{item['id']} {side}: answerable 값 오류")
    if data.get("route") not in ROUTES:
        errors.append(f"{item['id']} {side}: route 값 오류 {data.get('route')}")
    if not data.get("answer"):
        errors.append(f"{item['id']} {side}: answer 없음")
    if data["answerable"] in ("no", "partial") and not data.get("route"):
        errors.append(f"{item['id']} {side}: 답을 못 하는데 route 없음")
    for ev in data.get("evidence", []):
        did, _, page = ev.partition(":")
        if did not in pages:
            errors.append(f"{item['id']} {side}: 없는 문서 {did}")
        elif side == "v1" and did.startswith("notice"):
            errors.append(f"{item['id']} v1: 공지를 근거로 씀 {did}")
        elif page and page not in pages[did]:
            errors.append(f"{item['id']} {side}: {did}에 페이지 {page} 청크 없음")


def export_xlsx(items, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "평가셋"
    head = ["ID", "유형", "멘토링", "출처", "질문",
            "v1 답변가능", "v1 정답 (매뉴얼만)", "v1 근거", "v1 문의처",
            "v2 답변가능", "v2 정답 (공지 포함)", "v2 근거", "v2 문의처", "메모",
            "검수: 실제 말투인가 (O/X)", "검수: 정답 맞나 (O/X)", "검수 의견"]
    ws.append(head)
    for it in items:
        v1, v2 = it["v1"], it["v2"]
        note = " / ".join(x for x in [v1.get("note"), (v2 or {}).get("note")] if x)
        ws.append([
            it["id"], it["category"], it["type"], it["origin"], it["question"],
            v1["answerable"], v1["answer"], "\n".join(v1["evidence"]), v1["route"] or "",
            (v2 or v1)["answerable"], "(v1과 같음)" if v2 is None else v2["answer"],
            "" if v2 is None else "\n".join(v2["evidence"]), "" if v2 is None else (v2["route"] or ""),
            note, "", "", "",
        ])
    widths = [7, 11, 8, 8, 34, 9, 52, 22, 12, 9, 52, 22, 12, 24, 12, 12, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    fill = PatternFill("solid", fgColor="E8EEF7")
    review = PatternFill("solid", fgColor="FFF4D6")
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = review if c.column >= 15 else fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"O,X"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"O2:P{len(items) + 1}")
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main():
    data = json.loads((ROOT / "eval/eval_set.json").read_text(encoding="utf-8"))
    items, pages, errors = data["items"], chunk_pages(), []
    ids = [it["id"] for it in items]
    if len(ids) != len(set(ids)):
        errors.append("ID 중복")
    for it in items:
        check_side(it, "v1", it["v1"], pages, errors)
        if it["v2"] is not None:
            check_side(it, "v2", it["v2"], pages, errors)
    print(f"문항 {len(items)}개")
    for key in ("category", "origin", "type"):
        print(f"  {key}:", dict(Counter(it[key] for it in items)))
    print("  v1 answerable:", dict(Counter(it["v1"]["answerable"] for it in items)))
    print("  v2 answerable:", dict(Counter((it["v2"] or it["v1"])["answerable"] for it in items)))
    print("  v1과 v2 정답이 다른 문항:", sum(1 for it in items if it["v2"] is not None))
    for e in errors:
        print("ERROR", e)
    if not errors:
        export_xlsx(items, ROOT / "eval/eval_set_review.xlsx")
        print("OK, eval/eval_set_review.xlsx 생성")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
